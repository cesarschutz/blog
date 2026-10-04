"""O estilo do blog nas apresentações (D74, skill `apresentacao`).

Biblioteca para montar o .pptx de um post com a identidade do site: o papel e a tinta, a ficha de
catálogo, a folha com sombra curta, a fita, o post-it, a ficha pautada, o bloco de código do site e a
caneta do caderno (as marcações da D48), com as fontes do site e os desenhos do próprio post.

As cores vêm de `src/styles/tokens.ts` (claro, escuro, os tons dos diagramas e o marca-texto) e a cor
do livro, de `src/livros/livros.json`: se o visual do site mudar, as apresentações acompanham.

O texto é medido com as fontes (Pillow) e quebrado aqui, linha por linha, para cada marcação cair no
lugar certo. Com o espaçamento exato, a linha de base fica a 0,205 em do pé da linha, em qualquer fonte
(medido no LibreOffice em 02/10/2026; no PowerPoint e no Keynote as marcações podem andar um pouco).

Uso: o roteiro de cada post fica em `scripts/slides/posts/<slug>.py` (o exemplo é o do post dos mods).
"""
import json
import math
import random
import re
import shutil
import uuid
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

RAIZ = Path(__file__).resolve().parents[2]
SAIDA = RAIZ / "saida" / "slides"
SITE = "blog.cesarschutz.com.br"

# O quadro: 16:9 (13,333 × 7,5 polegadas), margens de 0,75.
W, H = 13.333, 7.5
ML, MR = 0.75, 0.75
LARG = W - ML - MR


# ---------------------------------------------------------------- cores (src/styles/tokens.ts)
def _objeto(texto, nome):
    i = texto.index(f"export const {nome}")
    j = texto.index("{", i)
    prof = 0
    for k in range(j, len(texto)):
        if texto[k] == "{":
            prof += 1
        elif texto[k] == "}":
            prof -= 1
            if prof == 0:
                return texto[j + 1:k]
    raise ValueError(nome)


def _cores(corpo):
    return {(m.group(1) or m.group(2)): m.group(3).upper()
            for m in re.finditer(r'(?:"([\w-]+)"|(\w+)):\s*"#([0-9A-Fa-f]{6})"', corpo)}


_TOKENS = (RAIZ / "src/styles/tokens.ts").read_text()
CLARO = _cores(_objeto(_TOKENS, "claro"))
ESCURO = _cores(_objeto(_TOKENS, "escuro"))
TONS = _cores(_objeto(_TOKENS, "DIAGRAMA"))
_MT = re.search(r'MARCA_TEXTO\s*=\s*\{\s*claro:\s*\{\s*cor:\s*"#([0-9A-Fa-f]{6})"', _TOKENS)


def _canais(h):
    return [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]


def _lin(v):
    return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4


def _gama(v):
    return 12.92 * v if v <= 0.0031308 else 1.055 * v ** (1 / 2.4) - 0.055


def _oklab(h):
    r, g, b = [_lin(v) for v in _canais(h)]
    l_ = (0.4122214708 * r + 0.5363325363 * g + 0.0514459929 * b) ** (1 / 3)
    m_ = (0.2119034982 * r + 0.6806995451 * g + 0.1073969566 * b) ** (1 / 3)
    s_ = (0.0883024619 * r + 0.2817188376 * g + 0.6299787005 * b) ** (1 / 3)
    return [0.2104542553 * l_ + 0.793617785 * m_ - 0.0040720468 * s_,
            1.9779984951 * l_ - 2.428592205 * m_ + 0.4505937099 * s_,
            0.0259040371 * l_ + 0.7827717662 * m_ - 0.808675766 * s_]


def _de_oklab(lab):
    L, A, B = lab
    l_ = (L + 0.3963377774 * A + 0.2158037573 * B) ** 3
    m_ = (L - 0.1055613458 * A - 0.0638541728 * B) ** 3
    s_ = (L - 0.0894841775 * A - 1.291485548 * B) ** 3
    rgb = [4.0767416621 * l_ - 3.3077115913 * m_ + 0.2309699292 * s_,
           -1.2684380046 * l_ + 2.6097574011 * m_ - 0.3413193965 * s_,
           -0.0041960863 * l_ - 0.7034186147 * m_ + 1.707614701 * s_]
    return "".join(f"{round(min(1, max(0, _gama(v))) * 255):02X}" for v in rgb)


def misturar(a, b, p):
    """Igual a `color-mix(in oklab, a, b p%)` do site (misturar() de tokens.ts)."""
    x, y = _oklab(a), _oklab(b)
    return _de_oklab([x[i] * (1 - p / 100) + y[i] * (p / 100) for i in range(3)])


PAPER, FOLHA, WELL = CLARO["paper"], CLARO["paper-hi"], CLARO["well"]
INK, INK2, INK3, RULE = CLARO["ink"], CLARO["ink-2"], CLARO["ink-3"], CLARO["rule"]
ACENTO, CANETA = CLARO["acento"], CLARO["caneta"]
POSTIT, FITA = CLARO["post-it"], CLARO["fita"]
SOMBRA_QUENTE = CLARO.get("sombra-quente", "3C2608")
NOTA, DICA, IMPORTANTE = CLARO["aviso-nota"], CLARO["aviso-dica"], CLARO["aviso-importante"]
ATENCAO, CUIDADO = CLARO["aviso-atencao"], CLARO["aviso-cuidado"]
E_PAPER, E_FOLHA, E_INK, E_INK2 = ESCURO["paper"], ESCURO["paper-hi"], ESCURO["ink"], ESCURO["ink-2"]
E_INK3, E_RULE, E_CANETA = ESCURO["ink-3"], ESCURO["rule"], ESCURO["caneta"]
AZUL, VERDE, AMBAR = TONS["azul"], TONS["verde"], TONS["ambar"]
VERMELHO, ROXO, PETROLEO = TONS["vermelho"], TONS["roxo"], TONS["petroleo"]
MARCA_TEXTO = _MT.group(1).upper() if _MT else "FFE27A"
# A barra do bloco de código (src/lib/codigo.ts), a pauta e o cabeçalho da ficha pautada (C05).
BARRA = misturar(FOLHA, INK, 8)
PAUTA = misturar(FOLHA, CANETA, 18)
CABECALHO = misturar(FOLHA, CUIDADO, 50)
BRANCO = "FFFFFF"

SOMBRA = dict(blur=4, dist=1.5, dir=90, cor=SOMBRA_QUENTE, alpha=0.22)
SOMBRA_POSTIT = dict(blur=9, dist=3.5, dir=100, cor=SOMBRA_QUENTE, alpha=0.24)
SOMBRA_FOTO = dict(blur=5, dist=2, dir=90, cor=SOMBRA_QUENTE, alpha=0.2)


def livro(categoria):
    """O livro do post (src/livros/livros.json): volume e cor, pela categoria."""
    dados = json.loads((RAIZ / "src/livros/livros.json").read_text())
    for lv in dados["livros"]:
        if categoria in (lv.get("titulo"), lv.get("slug")):
            return dict(volume=lv["volume"], cor=lv["cor"].lstrip("#").upper(), titulo=lv["titulo"])
    return dict(volume=None, cor=INK2, titulo=categoria)


def ler_post(slug):
    """O frontmatter do post (o essencial) e o corpo, de src/content/posts/<slug>.md ou .mdx."""
    for ext in (".mdx", ".md"):
        p = RAIZ / "src/content/posts" / f"{slug}{ext}"
        if p.exists():
            texto = p.read_text()
            break
    else:
        raise SystemExit(f"Post não encontrado: {slug}")
    _, fm, corpo = texto.split("---", 2)
    dados = {}
    for chave in ("title", "description", "published", "category", "formato"):
        m = re.search(rf'^{chave}:\s*"?(.*?)"?\s*$', fm, re.M)
        if m:
            dados[chave] = m.group(1)
    m = re.search(r"^tags:\s*\[(.*?)\]", fm, re.M)
    dados["tags"] = [t.strip().strip('"') for t in m.group(1).split(",")] if m else []
    dados["corpo"] = corpo
    return dados


# ---------------------------------------------------------------- fontes do site (em TTF)
TITULO, TEXTO, UI, MONO, MAO = "Besley", "Literata", "IBM Plex Sans", "JetBrains Mono", "Caveat"
_ARQUIVO = {TITULO: "Besley", TEXTO: "Literata", UI: "IBMPlexSans", MONO: "JetBrainsMono", MAO: "Caveat"}
_PASTAS = [Path.home() / "Library/Fonts", Path("/Library/Fonts"), Path.home() / ".local/share/fonts",
           Path.home() / ".fonts", Path("/usr/share/fonts/truetype")]
_pil = {}


def arquivo_fonte(familia, negrito=False, italico=False):
    negrito, italico = bool(negrito), bool(italico)
    estilos = [{(0, 0): "Regular", (1, 0): "Bold", (0, 1): "Italic", (1, 1): "BoldItalic"}[(int(negrito), int(italico))]]
    if italico:  # a Caveat não tem itálico
        estilos.append("Bold" if negrito else "Regular")
    for estilo in estilos:
        for pasta in _PASTAS:
            p = pasta / f"{_ARQUIVO[familia]}-{estilo}.ttf"
            if p.exists():
                return p
    raise SystemExit(f"Falta a fonte {familia} ({estilos[0]}) em TTF. Instale com "
                     "`python3 scripts/slides/fontes.py`, que baixa do Google Fonts (pergunte ao Cesar antes).")


def pil(f, b=False, i=False):
    k = (f, bool(b), bool(i))
    if k not in _pil:
        _pil[k] = ImageFont.truetype(str(arquivo_fonte(f, b, i)), 1000)
    return _pil[k]


def larg(t, f, s, b=False, i=False):
    """A largura do texto, em polegadas."""
    return pil(f, b, i).getlength(t) / 1000 * s / 72


def larg_run(r):
    return larg(r["t"], r["f"], r["s"], r.get("b"), r.get("i"))


# ---------------------------------------------------------------- texto com marcações
TOK = re.compile(r"\{(\w+):(.+?)\}|`([^`]+)`|\*\*(.+?)\*\*|~(.+?)~")


def analisar(texto, base, marca=None):
    """`código`, **negrito**, ~itálico~ e {id:trecho} (o trecho que a caneta marca); \\n quebra a linha."""
    runs, pos = [], 0

    def push(t, est):
        if t:
            r = {**est, "t": t}
            if marca:
                r["marca"] = marca
            runs.append(r)

    for m in TOK.finditer(texto):
        push(texto[pos:m.start()], base)
        if m.group(1):
            runs.extend(analisar(m.group(2), base, m.group(1)))
        elif m.group(3) is not None:
            cod = {**base, "f": MONO, "s": round(base["s"] * 0.88, 1), "b": False, "i": False}
            cod.update(base.get("codigo", {}))
            push(m.group(3), cod)
        elif m.group(4) is not None:
            runs.extend(analisar(m.group(4), {**base, "b": True}, marca))
        elif m.group(5) is not None:
            runs.extend(analisar(m.group(5), {**base, "i": True}, marca))
        pos = m.end()
    push(texto[pos:], base)
    return runs


def dividir_quebras(runs):
    segs = [[]]
    for r in runs:
        for k, p in enumerate(r["t"].split("\n")):
            if k:
                segs.append([])
            if p:
                segs[-1].append({**r, "t": p})
    return segs


def quebrar(runs, w, folga=0.95):
    """Quebra as linhas com as larguras da fonte, para saber onde cai cada trecho marcado."""
    atomos = [{**r, "t": p} for r in runs for p in re.findall(r"\S+\s*|\s+", r["t"])]
    palavras, cur = [], []
    for a in atomos:
        cur.append(a)
        if a["t"][-1].isspace():
            palavras.append(cur)
            cur = []
    if cur:
        palavras.append(cur)
    linhas, atual = [[]], 0.0
    for pal in palavras:
        ult = pal[-1]
        lw = sum(larg(a["t"].rstrip() if k == len(pal) - 1 else a["t"], a["f"], a["s"], a.get("b"), a.get("i"))
                 for k, a in enumerate(pal))
        sp = larg_run(ult) - larg(ult["t"].rstrip(), ult["f"], ult["s"], ult.get("b"), ult.get("i"))
        if linhas[-1] and atual + lw > w * folga:
            linhas[-1][-1]["t"] = linhas[-1][-1]["t"].rstrip()
            linhas.append([])
            atual = 0.0
        linhas[-1].extend(dict(a) for a in pal)
        atual += lw + sp
    if linhas[-1]:
        linhas[-1][-1]["t"] = linhas[-1][-1]["t"].rstrip()
    return [l for l in linhas if l] or [[]]


class Diagrama:
    """Um bloco de texto já quebrado: as linhas, a altura e onde caiu cada trecho marcado."""


def diagramar(x, y, w, paragrafos, base, pitch=None, depois=0, alinhar="l", folga=0.95):
    d = Diagrama()
    d.x, d.y, d.w, d.pars, d.marcas = x, y, w, [], {}
    cy = y
    for par in paragrafos:
        if isinstance(par, str):
            par = {"t": par}
        est = {**base, **par.get("estilo", {})}
        pp = par.get("pitch", pitch) or round(est["s"] * 1.3, 1)
        dep = par.get("depois", depois)
        al = par.get("alinhar", alinhar)
        linhas = []
        for seg in dividir_quebras(analisar(par["t"], est)):
            linhas += quebrar(seg, w, folga) if seg else [[]]
        pin = pp / 72
        for k, linha in enumerate(linhas):
            lw = sum(larg_run(r) for r in linha)
            x0 = x if al == "l" else (x + (w - lw) / 2 if al == "c" else x + w - lw)
            topo = cy + k * pin
            base_y = topo + pin - 0.205 * max([r["s"] for r in linha] or [est["s"]]) / 72
            cx = x0
            for r in linha:
                rw = larg_run(r)
                if r.get("marca"):
                    lst = d.marcas.setdefault(r["marca"], [])
                    if lst and abs(lst[-1]["base"] - base_y) < 1e-6 and abs(lst[-1]["x1"] - cx) < 1e-6:
                        lst[-1]["x1"] = cx + rw
                    else:
                        lst.append(dict(x0=cx, x1=cx + rw, base=base_y, topo=topo, pitch=pin, tam=r["s"]))
                cx += rw
        d.pars.append(dict(linhas=linhas, pitch=pp, depois=dep, alinhar=al))
        cy += len(linhas) * pin + dep / 72
    d.altura = cy - y - (d.pars[-1]["depois"] / 72 if d.pars else 0)
    d.fim = y + d.altura
    return d


def formatar(run, r):
    f = run.font
    f.name = r["f"]
    f.size = Pt(r["s"])
    f.bold = bool(r.get("b"))
    f.italic = bool(r.get("i"))
    f.color.rgb = RGBColor.from_string(r["c"])
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "pt-BR")
    if r["f"] == MONO:
        rpr.set("noProof", "1")


def colocar(slide, d, shape=None, h=None, ancora="t", nome=None):
    """Põe o bloco diagramado numa caixa de texto (ou no título do slide)."""
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(d.x), Inches(d.y), Inches(d.w), Inches(h or d.altura + 0.06))
    else:
        shape.left, shape.top = Inches(d.x), Inches(d.y)
        shape.width, shape.height = Inches(d.w), Inches(h or d.altura + 0.06)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[ancora]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    primeiro = True
    for par in d.pars:
        p = tf.paragraphs[0] if primeiro else tf.add_paragraph()
        if primeiro:
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
        primeiro = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[par["alinhar"]]
        p.line_spacing = Pt(par["pitch"])
        p.space_after = Pt(par["depois"])
        p.space_before = Pt(0)
        for k, linha in enumerate(par["linhas"]):
            if k:
                p.add_line_break()
            for r in linha:
                run = p.add_run()
                run.text = r["t"]
                formatar(run, r)
    if nome:
        shape.name = nome
    return shape


# ---------------------------------------------------------------- formas
def sem_estilo(sh):
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)


def _cor(pai, cor, alpha=None):
    c = etree.SubElement(pai, qn("a:srgbClr"), val=cor)
    if alpha is not None:
        etree.SubElement(c, qn("a:alpha"), val=str(int(alpha * 100000)))
    return c


def _sombra(spPr, sombra):
    ef = etree.SubElement(spPr, qn("a:effectLst"))
    sd = etree.SubElement(ef, qn("a:outerShdw"), blurRad=str(int(sombra["blur"] * 12700)),
                          dist=str(int(sombra["dist"] * 12700)), dir=str(int(sombra["dir"] * 60000)),
                          algn="t", rotWithShape="0")
    _cor(sd, sombra["cor"], sombra["alpha"])


def pintar(sh, fundo=None, alpha=None, linha=None, lw=0.75, alpha_linha=None, tracejado=None, sombra=None):
    """Preenchimento, linha e sombra, na ordem que o PowerPoint exige (fill, ln, effectLst)."""
    spPr = sh._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill", "a:ln",
                "a:effectLst", "a:effectDag"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    if fundo is None:
        etree.SubElement(spPr, qn("a:noFill"))
    else:
        _cor(etree.SubElement(spPr, qn("a:solidFill")), fundo, alpha)
    ln = etree.SubElement(spPr, qn("a:ln"))
    if linha is None:
        etree.SubElement(ln, qn("a:noFill"))
    else:
        ln.set("w", str(int(lw * 12700)))
        ln.set("cap", "rnd")
        _cor(etree.SubElement(ln, qn("a:solidFill")), linha, alpha_linha)
        if tracejado:
            etree.SubElement(ln, qn("a:prstDash"), val=tracejado)
        etree.SubElement(ln, qn("a:round"))
    if sombra:
        _sombra(spPr, sombra)
    return sh


def retangulo(slide, x, y, w, h, raio=0.0, rot=0, tipo=None, nome=None, **pintura):
    t = tipo or (MSO_SHAPE.ROUNDED_RECTANGLE if raio else MSO_SHAPE.RECTANGLE)
    sh = slide.shapes.add_shape(t, Inches(x), Inches(y), Inches(w), Inches(h))
    sem_estilo(sh)
    if raio:
        sh.adjustments[0] = min(0.5, raio / min(w, h))
    if rot:
        sh.rotation = rot
    pintar(sh, **pintura)
    if nome:
        sh.name = nome
    return sh


def elipse(slide, cx, cy, r, **pintura):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    sem_estilo(sh)
    pintar(sh, **pintura)
    return sh


def livre(slide, pts, fechar=False, nome=None, **pintura):
    """Um traço livre (pontos em polegadas): a base de toda marca à caneta."""
    emu = [(int(px * 914400), int(py * 914400)) for px, py in pts]
    fb = slide.shapes.build_freeform(emu[0][0], emu[0][1], scale=1.0)
    fb.add_line_segments(emu[1:], close=fechar)
    sh = fb.convert_to_shape()
    sem_estilo(sh)
    pintar(sh, **pintura)
    if nome:
        sh.name = nome
    return sh


def por_tras(sh, ref):
    ref._element.addprevious(sh._element)


def fundo(slide, cor):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = RGBColor.from_string(cor)


def folha(slide, x, y, w, h, cor=FOLHA, borda=RULE, raio=0.07, sombra=True, nome=None):
    """A folha do site: papel claro, borda fina e a sombra curta e quente."""
    return retangulo(slide, x, y, w, h, raio=raio, fundo=cor, linha=borda, lw=0.75,
                     sombra=SOMBRA if sombra else None, nome=nome)


def fita(slide, cx, cy, w=1.05, h=0.3, rot=-4, alpha=0.74):
    """A fita adesiva translúcida que cola papéis e fotos."""
    return retangulo(slide, cx - w / 2, cy - h / 2, w, h, rot=rot, fundo=FITA, alpha=alpha, nome="Fita")


def imagem(slide, caminho, x, y, w=None, h=None, raio_px=None, alt="", recorte=None, rot=0, sombra=None):
    """Uma imagem com os cantos arredondados do painel (raio em pixels da imagem) e o texto alternativo."""
    caminho = Path(caminho)
    pw, ph = Image.open(caminho).size
    l, t, r, b = recorte or (0, 0, 0, 0)
    ew, eh = pw - l - r, ph - t - b
    if w and not h:
        h = w * eh / ew
    if h and not w:
        w = h * ew / eh
    pic = slide.shapes.add_picture(str(caminho), Inches(x), Inches(y), Inches(w), Inches(h))
    if recorte:
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l / pw, t / ph, r / pw, b / ph
    if raio_px:
        raio = raio_px * w / ew
        prst = pic._element.spPr.find(qn("a:prstGeom"))
        prst.set("prst", "roundRect")
        av = prst.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(prst, qn("a:avLst"))
        etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {int(min(0.5, raio / min(w, h)) * 100000)}")
    if rot:
        pic.rotation = rot
    if sombra:
        _sombra(pic._element.spPr, sombra)
    pic._element.nvPicPr.cNvPr.set("descr", alt)
    pic.name = caminho.stem
    return pic, w, h


# ---------------------------------------------------------------- a caneta do caderno (D48)
def _jit(rnd, a):
    return rnd.uniform(-a, a)


def tremido(pts, seed=0, a=0.004, passo=0.18):
    """Subdivide e treme de leve um traço reto, como à mão."""
    rnd = random.Random(seed)
    out = [pts[0]]
    for (xa, ya), (xb, yb) in zip(pts, pts[1:]):
        n = max(1, int(math.hypot(xb - xa, yb - ya) / passo))
        for k in range(1, n + 1):
            t = k / n
            j = _jit(rnd, a) if k < n else 0
            out.append((xa + (xb - xa) * t + j, ya + (yb - ya) * t + j))
    return out


def curva(p0, p1, p2, n=24):
    """Uma curva de Bézier quadrática (para as setas)."""
    return [((1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * p1[0] + t * t * p2[0],
             (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * p1[1] + t * t * p2[1]) for t in (i / n for i in range(n + 1))]


def seta(slide, pts, cor=CANETA, lw=1.5, ponta=0.09):
    (xa, ya), (xb, yb) = pts[-2], pts[-1]
    ang = math.atan2(yb - ya, xb - xa)
    w1 = (xb - ponta * math.cos(ang - 0.5), yb - ponta * math.sin(ang - 0.5))
    w2 = (xb - ponta * math.cos(ang + 0.5), yb - ponta * math.sin(ang + 0.5))
    return livre(slide, list(pts) + [w1, (xb, yb), w2], linha=cor, lw=lw, nome="Seta à caneta")


def ondulado(slide, seg, cor=CANETA, lw=1.4):
    k = seg["tam"] / 16
    y = seg["base"] + 0.062 * k
    x0, x1 = seg["x0"] - 0.01, seg["x1"] + 0.01
    amp, comp = 0.021 * k, 0.125 * k
    n = max(10, int((x1 - x0) / 0.012))
    pts = [(x0 + (x1 - x0) * i / n, y + amp * math.sin(2 * math.pi * ((x1 - x0) * i / n) / comp)) for i in range(n + 1)]
    return livre(slide, pts, linha=cor, lw=lw, nome="Sublinhado ondulado")


def duplo(slide, seg, cor=CANETA, lw=1.3, seed=3):
    k = seg["tam"] / 16
    rnd = random.Random(seed)
    y1 = seg["base"] + 0.05 * k
    y2 = y1 + 0.052 * k
    x0, x1 = seg["x0"], seg["x1"]
    a = [(x0 - 0.02 + (x1 - x0 + 0.05) * i / 12, y1 + 0.008 * math.sin(math.pi * i / 12) + _jit(rnd, 0.002)) for i in range(13)]
    b = [(x0 + 0.03 + (x1 - x0 - 0.02) * i / 12, y2 + 0.006 * math.sin(math.pi * i / 12) + _jit(rnd, 0.002)) for i in range(13)]
    livre(slide, a, linha=cor, lw=lw, nome="Sublinhado duplo")
    livre(slide, b, linha=cor, lw=lw, nome="Sublinhado duplo")


def sublinhado(slide, x0, x1, y, cor=CANETA, lw=1.4, seed=9):
    return livre(slide, tremido([(x0, y), (x1, y - 0.005)], seed, 0.003, 0.15), linha=cor, lw=lw, nome="Sublinhado à caneta")


def laco(cx, cy, rx, ry, seed=1, inicio=208, giro=382, incl=-2.5):
    """Um círculo à mão que se cruza no alto, à esquerda."""
    rnd = random.Random(seed)
    pts, a = [], math.radians(incl)
    n = 80
    for i in range(n + 1):
        t = i / n
        ang = math.radians(inicio + giro * t)
        k = (1.02 - 0.035 * t) * (1 + 0.015 * math.sin(2 * ang + seed)) + _jit(rnd, 0.003)
        px, py = rx * k * math.cos(ang), ry * k * math.sin(ang)
        pts.append((cx + px * math.cos(a) - py * math.sin(a), cy + px * math.sin(a) + py * math.cos(a)))
    return pts


def circulo(slide, seg=None, cx=None, cy=None, rx=None, ry=None, cor=CANETA, lw=1.6, seed=1):
    """O círculo em volta de um trecho (seg) ou de um ponto de uma imagem (cx, cy, rx, ry)."""
    if seg:
        k = seg["tam"] / 72
        cx = (seg["x0"] + seg["x1"]) / 2
        cy = seg["base"] - 0.33 * k
        # a elipse passa por fora dos cantos do trecho: (a/rx)² + (b/ry)² <= 1
        a, b = (seg["x1"] - seg["x0"]) / 2, 0.36 * k
        ry = 0.64 * k + 0.05
        rx = a / math.sqrt(max(0.3, 1 - (b / ry) ** 2)) + 0.02
    return livre(slide, laco(cx, cy, rx, ry, seed), linha=cor, lw=lw, nome="Círculo à caneta")


def caixa(slide, seg=None, x=None, y=None, w=None, h=None, cor=CANETA, lw=1.3, seed=2):
    if seg:
        k = seg["tam"] / 72
        x, w = seg["x0"] - 0.045, seg["x1"] - seg["x0"] + 0.075
        y, h = seg["base"] - 0.8 * k, 1.06 * k
    rnd = random.Random(seed)
    j = lambda: _jit(rnd, 0.008)  # noqa: E731
    pts = [(x + 0.02, y + j()), (x + w + j(), y + j()), (x + w + j(), y + h + j()), (x + j(), y + h + j()),
           (x + j(), y + j()), (x + 0.09, y + 0.005)]
    return livre(slide, tremido(pts, seed, 0.003, 0.12), linha=cor, lw=lw, nome="Caixa à caneta")


def faixa(slide, seg, baixa=False, seed=0, alpha=0.52):
    """O marca-texto (no máximo três por apresentação) e o marca-texto baixo, o grifo."""
    k = seg["tam"] / 72
    rnd = random.Random(seed)
    if baixa:
        topo, pe = seg["base"] - 0.34 * k, seg["base"] + 0.12 * k
    else:
        topo, pe = seg["base"] - 0.80 * k, seg["base"] + 0.20 * k
    x0, x1 = seg["x0"] - 0.035, seg["x1"] + 0.035
    pts = [(x0 + 0.01, topo + _jit(rnd, 0.01)), ((x0 + x1) / 2, topo - 0.006), (x1, topo + _jit(rnd, 0.01)),
           (x1 + 0.012, (topo + pe) / 2), (x1 - 0.004, pe + _jit(rnd, 0.008)), ((x0 + x1) / 2, pe + 0.006),
           (x0, pe + _jit(rnd, 0.008)), (x0 - 0.012, (topo + pe) / 2)]
    return livre(slide, pts, fechar=True, fundo=MARCA_TEXTO, alpha=alpha, nome="Marca-texto")


def colchete(slide, x, y0, y1, cor=CANETA, lw=1.6, aba=0.13):
    pts = [(x + aba, y0), (x + 0.025, y0 + 0.012), (x, y0 + 0.06), (x - 0.004, (y0 + y1) / 2), (x, y1 - 0.06),
           (x + 0.025, y1 - 0.012), (x + aba, y1)]
    return livre(slide, tremido(pts, 7, 0.003, 0.2), linha=cor, lw=lw, nome="Colchete na margem")


def chave(slide, x, y0, y1, cor=CANETA, lw=1.5, r=0.075):
    """A chave que agrupa itens de lista, com a ponta para a direita (a nota vai ao lado)."""
    ym = (y0 + y1) / 2

    def arco(cx, cy, a0, a1, n=7):
        return [(cx + r * math.cos(math.radians(a0 + (a1 - a0) * i / n)),
                 cy + r * math.sin(math.radians(a0 + (a1 - a0) * i / n))) for i in range(n + 1)]
    pts = arco(x, y0 + r, -90, 0) + [(x + r, ym - r)] + arco(x + 2 * r, ym - r, 180, 90) + \
        arco(x + 2 * r, ym + r, 270, 180) + [(x + r, y1 - r)] + arco(x, y1 - r, 0, 90)
    return livre(slide, pts, linha=cor, lw=lw, nome="Chave agrupando")


def visto(slide, x, y, s=1.0, cor=CANETA, lw=2.0):
    pts = curva((x, y + 0.13 * s), (x + 0.05 * s, y + 0.19 * s), (x + 0.085 * s, y + 0.25 * s), 6) + \
        curva((x + 0.085 * s, y + 0.25 * s), (x + 0.15 * s, y + 0.06 * s), (x + 0.33 * s, y - 0.04 * s), 10)[1:]
    return livre(slide, pts, linha=cor, lw=lw, nome="Visto na margem")


def moldura(slide, x, y, w, h, cor=CANETA, lw=1.5, seed=5):
    rnd = random.Random(seed)
    pts = [(x + 0.05, y + _jit(rnd, 0.01)), (x + w + 0.02, y + _jit(rnd, 0.01)), (x + w + _jit(rnd, 0.01), y + h),
           (x + _jit(rnd, 0.01), y + h + 0.01), (x - 0.01, y - 0.02), (x + 0.2, y + 0.005)]
    return livre(slide, tremido(pts, seed, 0.005, 0.25), linha=cor, lw=lw, nome="Moldura à caneta")


def carimbo(slide, x, y, w, texto, h=0.6, rot=-3.5, cor=CANETA, tam=11.5):
    """O carimbo do que foi conferido ou rodado de verdade."""
    sh = retangulo(slide, x, y, w, h, raio=0.05, rot=rot, fundo=None, linha=cor, lw=1.6, nome="Carimbo")
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = texto
    formatar(run, dict(f=MONO, s=tam, c=cor, b=True))
    return sh


# ---------------------------------------------------------------- texto pronto e peças
def escrever(slide, x, y, w, pars, base, pitch=None, depois=0, alinhar="l", h=None, ancora="t", nome=None,
             marcas=None, caneta=CANETA, folga=0.95):
    """Escreve e marca: marcas={"id": "marca"|"grifo"|"ondulado"|"duplo"|"circulo"|"caixa"} para os {id:…}."""
    d = diagramar(x, y, w, pars, base, pitch, depois, alinhar, folga)
    marcas = marcas or {}
    for mid, tipo in marcas.items():
        if tipo in ("marca", "grifo"):
            for i, seg in enumerate(d.marcas.get(mid, [])):
                faixa(slide, seg, baixa=(tipo == "grifo"), seed=i)
    d.shape = colocar(slide, d, h=h, ancora=ancora, nome=nome)
    for mid, tipo in marcas.items():
        for i, seg in enumerate(d.marcas.get(mid, [])):
            if tipo == "ondulado":
                ondulado(slide, seg, cor=caneta)
            elif tipo == "duplo":
                duplo(slide, seg, cor=caneta, seed=i + 3)
            elif tipo == "circulo":
                circulo(slide, seg, cor=caneta, seed=i + 1)
            elif tipo == "caixa":
                caixa(slide, seg, cor=caneta, seed=i + 2)
    return d


def nota(slide, x, y, w, texto, tam=20, cor=CANETA, alinhar="l", pitch=None):
    """A nota à mão (Caveat), sempre na cor da caneta."""
    return escrever(slide, x, y, w, [texto], dict(f=MAO, s=tam, c=cor, b=True), pitch=pitch or tam * 1.05,
                    alinhar=alinhar, nome="Nota à mão", folga=0.98)


def ficha(slide, x, y, w, h, esq, dir_="", cor=INK2, tam=10, nome=None):
    """A ficha de catálogo: a tira em mono, sobre um fio de 2 pt na cor do livro (ou do papel da ficha)."""
    folha(slide, x, y, w, h, nome=nome or "Ficha")
    base = dict(f=MONO, s=tam, c=INK2)
    wd = larg(dir_, MONO, tam) + 0.05 if dir_ else 0
    escrever(slide, x + 0.17, y + 0.12, w - 0.34 - wd - 0.1, [esq], base, pitch=tam * 1.3, folga=1.0)
    if dir_:
        escrever(slide, x + w - 0.17 - wd, y + 0.12, wd, [dir_], base, pitch=tam * 1.3, alinhar="r", folga=1.0)
    retangulo(slide, x, y + 0.39, w, 0.028, fundo=cor, nome="Fio da ficha")
    return y + 0.39 + 0.028


def ficha_pautada(slide, x, y, w, h, titulo, pautas, tam_titulo=25):
    """A ficha pautada da papelaria (C05): cabeçalho à mão, linha vermelha e pautas azuis."""
    folha(slide, x, y, w, h, nome="Ficha pautada")
    escrever(slide, x + 0.26, y + 0.08, w - 0.4, [titulo], dict(f=MAO, s=tam_titulo, c=INK, b=True),
             pitch=tam_titulo * 1.15)
    retangulo(slide, x, y + 0.62, w, 0.021, fundo=CABECALHO, nome="Linha do cabeçalho")
    for py in pautas:
        retangulo(slide, x, py, w, 0.016, fundo=PAUTA, nome="Pauta")
    return y + 0.62


def postit(slide, x, y, w, h, texto, rot=2.0, tam=21, cor_texto=INK):
    """O post-it: o código vai na cor da caneta; \\n quebra a linha."""
    sh = retangulo(slide, x, y, w, h, rot=rot, fundo=POSTIT, sombra=SOMBRA_POSTIT, nome="Post-it")
    tf = sh.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.2)
    tf.margin_top = tf.margin_bottom = Inches(0.14)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.line_spacing = Pt(tam * 1.08)
    runs = analisar(texto, dict(f=MAO, s=tam, c=cor_texto, b=False, codigo=dict(f=MAO, s=tam, c=CANETA, b=True)))
    for k, seg in enumerate(dividir_quebras(runs)):
        if k:
            p.add_line_break()
        for r in seg:
            run = p.add_run()
            run.text = r["t"]
            formatar(run, r)
    return sh


# As cores do código no site (src/lib/codigo.ts): chave e palavra reservada em azul-tinta, texto entre
# aspas no verde do Dica, tipo no azul do Nota, comentário em ink-3 itálico e pontuação em ink-2.
COD_CHAVE, COD_TEXTO, COD_TIPO = ACENTO, DICA, NOTA
COD_COMENTARIO, COD_PONTO, COD_NUMERO, COD_NOME = INK3, INK2, IMPORTANTE, INK


def codigo(slide, x, y, w, aba, linhas, tam=12, pitch=None, destaque=None, nome=None):
    """O bloco de código do site: a barra, a aba do arquivo e a superfície. Cada linha é uma lista de
    (texto, cor) ou (texto, cor, "i"|"b"); destaque = o índice da linha que leva o marca-texto."""
    pitch = pitch or round(tam * 1.55, 1)
    pin = pitch / 72
    hc, pad = 0.4, 0.17
    h = hc + 2 * pad + len(linhas) * pin
    retangulo(slide, x, y, w, h, raio=0.07, fundo=WELL, linha=RULE, lw=0.75, sombra=SOMBRA, nome=nome or "Código")
    barra = retangulo(slide, x, y, w, hc, tipo=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, fundo=BARRA, linha=None,
                      nome="Barra do código")
    barra.adjustments[0] = 0.07 / hc
    barra.adjustments[1] = 0
    retangulo(slide, x, y + hc - 0.01, w, 0.011, fundo=RULE, nome="Linha da barra")
    tw = larg(aba, MONO, 10.5) + 0.36
    retangulo(slide, x + 0.14, y + 0.09, tw, hc - 0.09 + 0.002, fundo=WELL, nome="Aba do arquivo")
    escrever(slide, x + 0.14 + 0.18, y + 0.17, tw, [aba], dict(f=MONO, s=10.5, c=INK2), pitch=13)
    topo = y + hc + pad
    if destaque is not None:
        retangulo(slide, x + 0.02, topo + destaque * pin, w - 0.04, pin, fundo=MARCA_TEXTO, alpha=0.42,
                  nome="Linha destacada")
    tb = slide.shapes.add_textbox(Inches(x + 0.3), Inches(topo), Inches(w - 0.4), Inches(len(linhas) * pin + 0.05))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, linha in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(pitch)
        p.space_after = Pt(0)
        p.alignment = PP_ALIGN.LEFT
        for txt, cor, *ext in linha:
            if not txt:
                continue
            sp = len(txt) - len(txt.lstrip(" "))
            run = p.add_run()
            run.text = " " * sp + txt[sp:]
            formatar(run, dict(f=MONO, s=tam, c=cor, i=bool(ext and ext[0] == "i"), b=bool(ext and ext[0] == "b")))
    tb.name = "Código: " + aba
    return dict(x=x, y=y, w=w, h=h, topo=topo, pin=pin, texto_x=x + 0.3, tam=tam)


def aviso(slide, x, y, w, rotulo, texto, cor=ATENCAO, h=0.98, tam=13):
    """O aviso do site (Atenção, Cuidado…): folha tingida, ícone e rótulo na cor do aviso."""
    retangulo(slide, x, y, w, h, raio=0.07, fundo=misturar(FOLHA, cor, 7), linha=cor, lw=0.75, alpha_linha=0.45,
              sombra=SOMBRA, nome=f"Aviso: {rotulo}")
    tri = [(x + 0.36, y + 0.24), (x + 0.53, y + 0.53), (x + 0.19, y + 0.53)]
    livre(slide, tri + [tri[0]], linha=cor, lw=1.5, nome="Ícone do aviso")
    escrever(slide, x + 0.3, y + 0.3, 0.12, ["!"], dict(f=UI, s=11, c=cor, b=True), pitch=13, alinhar="c")
    escrever(slide, x + 0.72, y + 0.17, 2.5, [rotulo], dict(f=UI, s=13.5, c=cor, b=True), pitch=17)
    return escrever(slide, x + 0.72, y + 0.46, w - 1.0, [texto], dict(f=TEXTO, s=tam, c=INK), pitch=tam * 1.3)


def numero_grande(slide, x, y, w, valor, rotulo, cor=INK, tam=40):
    """Um dado em destaque: o número em Besley e o rótulo embaixo."""
    escrever(slide, x, y, w, [valor], dict(f=TITULO, s=tam, c=cor, b=True), pitch=tam * 1.15)
    return escrever(slide, x, y + tam * 1.15 / 72 + 0.08, w - 0.3, [rotulo], dict(f=UI, s=12.5, c=INK2), pitch=16)


def tabela(slide, x, y, w, colunas, linhas, altura_linha=0.6, cabecalho=True):
    """A tabela numa folha: colunas = [(título, largura, estilo)]. Cada célula é um texto (com as marcações
    de analisar()) ou uma lista de (texto, mudanças no estilo), numa linha só; fica no meio da linha."""
    h = (0.44 if cabecalho else 0.04) + len(linhas) * altura_linha + 0.04
    folha(slide, x, y, w, h, nome="Tabela")
    topo = y + 0.04
    if cabecalho:
        retangulo(slide, x + 0.01, y + 0.01, w - 0.02, 0.43, tipo=MSO_SHAPE.ROUND_2_SAME_RECTANGLE, fundo=WELL,
                  nome="Cabeçalho da tabela").adjustments[0] = 0.06 / 0.43
        cx = x + 0.25
        for titulo, cw, _ in colunas:
            escrever(slide, cx, y + 0.13, cw, [titulo], dict(f=UI, s=11.5, c=INK2, b=True), pitch=14)
            cx += cw
        topo = y + 0.44
    for k, cel in enumerate(linhas):
        ry = topo + k * altura_linha
        if k:
            retangulo(slide, x + 0.15, ry, w - 0.3, 0.011, fundo=RULE, nome="Linha da tabela")
        cx = x + 0.25
        for (titulo, cw, est), txt in zip(colunas, cel):
            if isinstance(txt, list):
                dd = diagramar(cx, 0, cw - 0.3, ["".join(t for t, _ in txt)], est, pitch=round(est["s"] * 1.24, 1))
                dd.pars[0]["linhas"] = [[{**est, **mud, "t": t} for t, mud in txt]]
            else:
                dd = diagramar(cx, 0, cw - 0.3, [txt], est, pitch=round(est["s"] * 1.24, 1))
            dd.y = ry + (altura_linha - dd.altura) / 2
            colocar(slide, dd)
            cx += cw
    return y + h


def escada(slide, itens, x0, y_baixo, largura, subida, destaque=PETROLEO, resposta=("não", "sim")):
    """A linha do tempo em escada: um degrau por item (nome, data, versão), do mais antigo ao mais novo,
    com a resposta à mão embaixo de cada degrau (o último é o destaque)."""
    n = len(itens)
    passo = largura / n
    pts = []
    for i in range(n):
        xa, yi = x0 + i * passo, y_baixo - i * subida
        if i == 0:
            pts.append((xa, yi))
        pts.append((xa + passo, yi))
        if i < n - 1:
            pts.append((xa + passo, yi - subida))
    livre(slide, tremido(pts, 11, 0.004, 0.2), linha=INK, lw=2.25, nome="A escada")
    livre(slide, [(x0 + (n - 1) * passo + 0.02, y_baixo - (n - 1) * subida), (x0 + n * passo, y_baixo - (n - 1) * subida)],
          linha=destaque, lw=4, nome="O último degrau")
    lw_max = passo - 0.1
    for i, (peca, data, versao) in enumerate(itens):
        xa, yi = x0 + i * passo, y_baixo - i * subida
        tam = 12.5
        while max(larg(p, UI, tam, True) for p in peca.split()) > lw_max and tam > 10:
            tam -= 0.25
        cor = destaque if i == n - 1 else INK
        pars = [{"t": peca, "estilo": dict(f=UI, s=tam, c=cor, b=True), "pitch": tam * 1.2, "depois": 3},
                {"t": data, "estilo": dict(f=MONO, s=9.5, c=INK2), "pitch": 12}]
        if versao:
            pars.append({"t": versao, "estilo": dict(f=MONO, s=9.5, c=INK3), "pitch": 12})
        d = diagramar(xa + 0.05, 0, lw_max, pars, dict(f=UI, s=tam, c=INK))
        d.y = yi - 0.1 - d.altura
        colocar(slide, d, nome=f"Degrau {i + 1}: {peca}")
        if resposta:
            ultimo = i == n - 1
            nota(slide, xa + passo - 0.5, yi + 0.03, 0.45, resposta[1] if ultimo else resposta[0], tam=17,
                 cor=destaque if ultimo else CANETA, alinhar="r")
    return passo


# ---------------------------------------------------------------- a apresentação
class Deck:
    """A apresentação de um post: slides 16:9 com o título no espaço reservado, as seções do PowerPoint,
    o número de cada slide, as notas do apresentador e o tema com as fontes e as cores do site."""

    def __init__(self, slug, titulo, assunto=""):
        self.slug, self.titulo, self.assunto = slug, titulo, assunto
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.layout = self.prs.slide_layouts[5]  # "Title Only": o título fica no espaço reservado
        self.secoes, self._numeros = [], []
        self.pasta = SAIDA / slug
        self.img = self.pasta / "img"
        self.post = ler_post(slug)
        self.livro = livro(self.post.get("category", ""))
        dados = self.pasta / "dados.json"
        self.tira = json.loads(dados.read_text()).get("tira") if dados.exists() else None

    def imagem(self, nome):
        """Uma foto do post (saida/slides/<slug>/img/, de capturar.mjs) ou um caminho do projeto."""
        p = self.img / nome
        return p if p.exists() else RAIZ / nome

    def slide(self, secao, escuro=False):
        s = self.prs.slides.add_slide(self.layout)
        fundo(s, E_PAPER if escuro else PAPER)
        if not self.secoes or self.secoes[-1][0] != secao:
            self.secoes.append((secao, []))
        self.secoes[-1][1].append(s.slide_id)
        return s

    def cabeca(self, s, secao, titulo, escuro=False, tam=34, y=0.8, w=None, marcas=None, caneta=None):
        """O título (no espaço reservado, para o sumário e a leitura), a seção do post e o número."""
        marcas = marcas or {}
        base = dict(f=TITULO, s=tam, c=E_INK if escuro else INK, b=True)
        d = diagramar(ML, y, w or LARG, [titulo], base, pitch=round(tam * 1.2, 1))
        t = s.shapes.title
        colocar(s, d, shape=t, nome="Título")
        for mid, tipo in marcas.items():
            for i, seg in enumerate(d.marcas.get(mid, [])):
                if tipo in ("marca", "grifo"):
                    por_tras(faixa(s, seg, baixa=(tipo == "grifo"), seed=i), t)
                elif tipo == "circulo":
                    circulo(s, seg, cor=caneta or (E_CANETA if escuro else CANETA), seed=i + 4, lw=1.9)
                elif tipo == "ondulado":
                    ondulado(s, seg, cor=caneta or (E_CANETA if escuro else CANETA), lw=1.8)
        ck = E_INK2 if escuro else INK2
        escrever(s, ML, 0.42, 9.0, [secao], dict(f=MONO, s=10.5, c=ck), pitch=13, nome="Seção do post")
        n = len(self.prs.slides)
        dn = escrever(s, W - MR - 2.0, 0.42, 2.0, [f"{n:02d} / 00"], dict(f=MONO, s=10.5, c=ck), pitch=13,
                      alinhar="r", nome="Número do slide")
        self._numeros.append((n, dn.shape))
        return d

    def capa(self, titulo, subtitulo, alt, secao="Abertura", rot=-0.6):
        """A abertura: a marca e as tags no alto, a tira da ficha, o título e o subtítulo do post, a data e
        o livro, e a capa do post colada com fita (as fotos de capturar.mjs)."""
        s = self.slide(secao)
        imagem(s, self.imagem("marca.png"), ML, 0.5, w=2.3, alt="Cesar Schutz, blog")
        x_tag, th = W - MR, 0.3
        for n in range(len(self.post["tags"]), 0, -1):
            arq = self.imagem(f"tag-{n}.png")
            if not arq.exists():
                continue
            pw, ph = Image.open(arq).size
            x_tag -= th * pw / ph
            imagem(s, arq, x_tag, 0.5, h=th, raio_px=31, alt=f"Tag {self.post['tags'][n - 1]}")
            x_tag -= 0.12
        tira = f"{self.tira['chamada']} · {self.tira['tombo']}" if self.tira else f"Vol. {self.livro['volume']:02d}"
        escrever(s, ML, 1.22, 6, [tira], dict(f=MONO, s=10.5, c=INK2), pitch=13, nome="Tira da ficha")
        data = "/".join(reversed(self.post["published"][:10].split("-")))
        escrever(s, W - MR - 5, 1.22, 5, [f"{data} · {self.livro['titulo']}"], dict(f=MONO, s=10.5, c=INK2),
                 pitch=13, alinhar="r")
        tam = 46
        while larg(titulo, TITULO, tam, True) > LARG * 0.95 and tam > 36:
            tam -= 1
        d = diagramar(ML, 1.55, LARG, [titulo], dict(f=TITULO, s=tam, c=INK, b=True), pitch=round(tam * 1.17, 1))
        colocar(s, d, shape=s.shapes.title, nome="Título")
        ds = escrever(s, ML, d.fim + 0.06, LARG, [subtitulo], dict(f=TEXTO, s=22, c=INK2, i=True), pitch=28,
                      nome="Subtítulo")
        arq = self.imagem("capa.png")
        if arq.exists():
            pw, ph = Image.open(arq).size
            topo, pe = ds.fim + 0.5, H - 0.4
            iw = min(10.9, (pe - topo) * (pw - 32) / (ph - 32))
            _, iw, ih = imagem(s, arq, (W - iw) / 2, topo, w=iw, raio_px=22, recorte=(16, 16, 16, 16), rot=rot,
                               sombra=SOMBRA_FOTO, alt=alt)
            fita(s, (W - iw) / 2 + 1.25, topo + 0.02, rot=-6)
            fita(s, (W + iw) / 2 - 1.25, topo - 0.04, rot=5)
        return s

    def fecho(self, titulo, notas, recomendacao=None, enderecos=(), fontes=(), secao_post="Fontes"):
        """O último slide: a recomendação final do post (com o visto), os endereços, as fontes resumidas, a
        ficha "Do livro" e a marca."""
        s = self.slide(self.secoes[-1][0] if self.secoes else "Fecho")
        self.cabeca(s, secao_post, titulo)
        y = 1.95
        if recomendacao:
            d = escrever(s, ML + 0.55, y, 7.3, [recomendacao], dict(f=TEXTO, s=19, c=INK), pitch=27)
            visto(s, ML, y + 0.05)
            y = d.fim + 0.25
        if enderecos:
            d = escrever(s, ML + 0.55, y, 7.3, list(enderecos), dict(f=MONO, s=12, c=ACENTO), pitch=18, depois=4)
            y = d.fim
        if fontes:
            escrever(s, ML, y + 0.55, 7.8, ["As fontes do post"], dict(f=UI, s=13, c=INK, b=True), pitch=16)
            escrever(s, ML, y + 0.9, 7.8, list(fontes), dict(f=UI, s=12, c=INK2), pitch=15.5, depois=5)
        livro_ = self.imagem("do-livro.png")
        if livro_.exists():
            imagem(s, livro_, W - MR - 2.75, 1.7, w=2.75,
                   alt=f"A ficha Do livro: o livro {self.livro['titulo']}, Volume {self.livro['volume']:02d}.")
        escrever(s, W - MR - 3.6, 6.12, 3.6, [SITE], dict(f=MONO, s=12, c=INK2), pitch=15, alinhar="r")
        imagem(s, self.imagem("marca.png"), W - MR - 2.3, 6.5, w=2.3, alt="Cesar Schutz, blog")
        self.notas(s, notas)
        return s

    @staticmethod
    def notas(s, texto):
        s.notes_slide.notes_text_frame.text = texto

    def _secoes_xml(self):
        el = self.prs.part._element
        ext_lst = el.find(qn("p:extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(el, qn("p:extLst"))
        p14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
        ext = etree.SubElement(ext_lst, qn("p:ext"), uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
        lst = etree.SubElement(ext, f"{{{p14}}}sectionLst", nsmap={"p14": p14})
        for nome, ids in self.secoes:
            sec = etree.SubElement(lst, f"{{{p14}}}section", name=nome, id="{" + str(uuid.uuid4()).upper() + "}")
            sl = etree.SubElement(sec, f"{{{p14}}}sldIdLst")
            for i in ids:
                etree.SubElement(sl, f"{{{p14}}}sldId", id=str(i))

    def salvar(self):
        total = len(self.prs.slides)
        for n, shape in self._numeros:
            runs = shape.text_frame.paragraphs[0].runs
            runs[0].text = f"{n:02d} / {total:02d}"
            for r in runs[1:]:  # a quebra separou "NN / 00" em pedaços: fica um só
                r._r.getparent().remove(r._r)
        self._secoes_xml()
        cp = self.prs.core_properties
        cp.title, cp.author, cp.language = self.titulo, "Cesar Schutz", "pt-BR"
        cp.subject = self.assunto or f"Apresentação do post do {SITE}"
        self.pasta.mkdir(parents=True, exist_ok=True)
        destino = self.pasta / f"{self.slug}.pptx"
        self.prs.save(destino)
        _ajustar_tema(destino)
        print("ok", destino)
        return destino


_TEMA_CORES = dict(dk1=INK, lt1=FOLHA, dk2=INK2, lt2=PAPER, accent1=ACENTO, accent2=PETROLEO, accent3=AZUL,
                   accent4=AMBAR, accent5=VERDE, accent6=ROXO, hlink=ACENTO, folHlink=IMPORTANTE)


def _ajustar_tema(caminho):
    """O tema do arquivo com as fontes (Besley e Literata) e as cores do site, para o que for criado depois."""
    tmp = caminho.with_suffix(".tmp")
    with zipfile.ZipFile(caminho) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            if re.match(r"ppt/theme/theme\d+\.xml$", item.filename):
                raiz = etree.fromstring(dados)
                te = raiz.find(qn("a:themeElements"))
                cs = te.find(qn("a:clrScheme"))
                cs.set("name", "Papel e luz")
                for nome_, val in _TEMA_CORES.items():
                    no = cs.find(qn(f"a:{nome_}"))
                    for ch in list(no):
                        no.remove(ch)
                    etree.SubElement(no, qn("a:srgbClr"), val=val)
                fs = te.find(qn("a:fontScheme"))
                fs.set("name", "Blog Cesar Schutz")
                fs.find(qn("a:majorFont")).find(qn("a:latin")).set("typeface", TITULO)
                fs.find(qn("a:minorFont")).find(qn("a:latin")).set("typeface", TEXTO)
                raiz.set("name", "Papel e luz")
                dados = etree.tostring(raiz, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, dados)
    shutil.move(tmp, caminho)

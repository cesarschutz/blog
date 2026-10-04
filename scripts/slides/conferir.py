"""A conferência da apresentação (D74, skill `apresentacao`): o PDF pelo LibreOffice, uma imagem por slide,
a folha de contato e as checagens que pegam erro antes de mostrar ao Cesar.

    python3 scripts/slides/conferir.py <slug> [--dpi 96]

Confere que todo slide tem título e notas do apresentador, que nenhuma forma sai do quadro, que o PDF só
usa as fontes do site (outra fonte é fonte que falta instalar) e a ordem do XML de cada forma
(preenchimento, linha, sombra), que o PowerPoint exige. Gera em saida/slides/<slug>/: o <slug>.pdf, as
imagens em previa/ e o contato.jpg (todos os slides numa folha). Sai com 1 se achar problema; olhar as
imagens, uma por uma, continua sendo obrigatório (texto que estoura, marcação fora do lugar).
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from estilo import SAIDA, H, W  # noqa: E402

FONTES_DO_SITE = ("Besley", "Literata", "IBMPlexSans", "JetBrainsMono", "Caveat")
ORDEM_SPPR = ["xfrm", "custGeom", "prstGeom", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
              "ln", "effectLst", "effectDag", "scene3d", "sp3d", "extLst"]
EMU = 914400


def soffice():
    for c in (shutil.which("soffice"), "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and Path(c).exists():
            return c
    sys.exit("Falta o LibreOffice (soffice) para gerar o PDF e as imagens dos slides.")


def main():
    args = sys.argv[1:]
    if not args:
        sys.exit(__doc__)
    slug = args[0]
    dpi = int(args[args.index("--dpi") + 1]) if "--dpi" in args else 96
    pasta = SAIDA / slug
    pptx = pasta / f"{slug}.pptx"
    if not pptx.exists():
        sys.exit(f"Não achei {pptx}: rode antes python3 scripts/slides/posts/{slug}.py")
    problemas = []

    prs = Presentation(pptx)
    for n, s in enumerate(prs.slides, 1):
        titulo = s.shapes.title
        if titulo is None or not titulo.text_frame.text.strip():
            problemas.append(f"slide {n}: sem título (o título vai no espaço reservado: deck.cabeca)")
        if not (s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()):
            problemas.append(f"slide {n}: sem notas do apresentador")
        for sh in s.shapes:
            if None in (sh.left, sh.top, sh.width, sh.height):
                continue
            x0, y0 = sh.left / EMU, sh.top / EMU
            x1, y1 = x0 + sh.width / EMU, y0 + sh.height / EMU
            if x0 < -0.02 or y0 < -0.02 or x1 > W + 0.02 or y1 > H + 0.02:
                problemas.append(f"slide {n}: '{sh.name}' sai do quadro ({x0:.2f}, {y0:.2f}) a ({x1:.2f}, {y1:.2f})")
        for spPr in s._element.iter(qn("p:spPr")):
            nomes = [el.tag.split("}")[1] for el in spPr]
            pos = [ORDEM_SPPR.index(t) for t in nomes if t in ORDEM_SPPR]
            if pos != sorted(pos):
                problemas.append(f"slide {n}: ordem do XML fora do esquema: {nomes}")

    with tempfile.TemporaryDirectory(prefix="lo_perfil_") as perfil:
        subprocess.run([soffice(), f"-env:UserInstallation=file://{perfil}", "--headless", "--convert-to", "pdf",
                        "--outdir", str(pasta), str(pptx)], check=True, capture_output=True, timeout=300)
    pdf = pasta / f"{slug}.pdf"
    fontes = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True, check=True).stdout.splitlines()[2:]
    usadas = sorted({linha.split()[0].split("+")[-1] for linha in fontes if linha.strip()})
    estranhas = [f for f in usadas if not f.startswith(FONTES_DO_SITE)]
    if estranhas:
        problemas.append(f"fontes fora do site no PDF (falta instalar ou falta o caractere): {', '.join(estranhas)}")

    previa = pasta / "previa"
    shutil.rmtree(previa, ignore_errors=True)
    previa.mkdir()
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), str(pdf), str(previa / "slide")], check=True)
    imagens = sorted(previa.glob("slide-*.png"))
    tw, th, cols, gap = 400, 225, 4, 12
    linhas = (len(imagens) + cols - 1) // cols
    folha = Image.new("RGB", (cols * tw + (cols + 1) * gap, linhas * th + (linhas + 1) * gap), (200, 200, 196))
    for i, f in enumerate(imagens):
        im = Image.open(f).convert("RGB").resize((tw, th), Image.LANCZOS)
        folha.paste(im, (gap + (i % cols) * (tw + gap), gap + (i // cols) * (th + gap)))
    folha.save(pasta / "contato.jpg", quality=88)

    print(f"{len(prs.slides)} slides · PDF: {pdf}")
    print(f"fontes: {', '.join(usadas)}")
    print(f"imagens: {previa}/ · folha de contato: {pasta / 'contato.jpg'}")
    if problemas:
        print("\nproblemas:")
        for p in problemas:
            print(" -", p)
        sys.exit(1)
    print("sem problemas nas checagens (falta olhar cada imagem)")


if __name__ == "__main__":
    main()
